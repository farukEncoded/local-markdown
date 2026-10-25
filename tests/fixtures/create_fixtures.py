"""Generate non-confidential, synthetic document fixtures for automated testing."""

from __future__ import annotations

import os
from pathlib import Path
import zipfile

import docx
import openpyxl
from pptx import Presentation
from PIL import Image, ImageDraw
import xlwt
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


def create_all_fixtures(output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)

    # 1. Synthetic DOCX
    docx_path = output_dir / "sample.docx"
    doc = docx.Document()
    doc.add_heading("Antigravity Project Overview", level=0)
    doc.add_paragraph("This is a synthetic test document containing critical business metrics.")
    doc.add_heading("Financial Summary", level=1)
    doc.add_paragraph("Total revenue reached $42,500 in Q3, representing a 14.5% year-over-year increase.")
    table = doc.add_table(rows=3, cols=3)
    table.rows[0].cells[0].text = "Quarter"
    table.rows[0].cells[1].text = "Revenue"
    table.rows[0].cells[2].text = "Margin"
    table.rows[1].cells[0].text = "Q2"
    table.rows[1].cells[1].text = "$37,100"
    table.rows[1].cells[2].text = "21.4%"
    table.rows[2].cells[0].text = "Q3"
    table.rows[2].cells[1].text = "$42,500"
    table.rows[2].cells[2].text = "23.8%"
    doc.add_paragraph("Documentation reference link: https://example.com/docs")
    doc.save(str(docx_path))

    # 2. Synthetic XLSX with visible and hidden sheets
    xlsx_path = output_dir / "sample.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Financials"
    ws1.append(["Department", "Allocation", "Spent"])
    ws1.append(["Engineering", 150000, 142000])
    ws1.append(["Design", 45000, 41200])
    ws1.append(["Marketing", 60000, 58900])

    ws2 = wb.create_sheet(title="Metadata")
    ws2.append(["Author", "Faruk Munshi"])
    ws2.append(["Version", "1.0.0"])

    ws3 = wb.create_sheet(title="ConfidentialData")
    ws3.append(["Code", "Target"])
    ws3.append(["PROJECT-X", 999999])
    ws3.sheet_state = "hidden"

    wb.save(str(xlsx_path))

    # 3. Synthetic genuine legacy XLS workbook
    xls_path = output_dir / "sample.xls"
    xls_book = xlwt.Workbook(encoding="utf-8")
    sheet1 = xls_book.add_sheet("Summary")
    sheet1.write(0, 0, "Item")
    sheet1.write(0, 1, "Amount")
    sheet1.write(1, 0, "Licenses")
    sheet1.write(1, 1, 1250)
    sheet1.write(2, 0, "Services")
    sheet1.write(2, 1, 3400)
    xls_book.save(str(xls_path))

    # 4. Synthetic PPTX
    pptx_path = output_dir / "sample.pptx"
    prs = Presentation()
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "Local Markdown Strategy"
    slide1.placeholders[1].text = "Client Offline Conversion Pipeline"

    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Key Deliverables"
    body = slide2.placeholders[1].text_frame
    body.text = "1. Zero network dependency during runtime."
    p2 = body.add_paragraph()
    p2.text = "2. Preserves all original source documents."
    prs.save(str(pptx_path))

    # 5. Synthetic PDF (pure text multi-page)
    pdf_path = output_dir / "sample.pdf"
    c = canvas.Canvas(str(pdf_path), pagesize=letter)
    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, 720, "Antigravity Executive Report - Page 1")
    c.setFont("Helvetica", 12)
    c.drawString(72, 690, "This page contains native digital text for verification.")
    c.drawString(72, 670, "Key metric: Operating margin increased to 28.5%.")
    c.showPage()

    c.setFont("Helvetica-Bold", 14)
    c.drawString(72, 720, "Financial Conclusion - Page 2")
    c.setFont("Helvetica", 12)
    c.drawString(72, 690, "All milestones have been verified successfully.")
    c.showPage()
    c.save()

    # 6. Synthetic corrupt document
    corrupt_path = output_dir / "corrupt.docx"
    corrupt_path.write_bytes(b"This is definitely not a valid PK zip archive.")

    # 7. Macro-enabled document
    macro_path = output_dir / "macro.docm"
    with zipfile.ZipFile(macro_path, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types></Types>")
        zf.writestr("word/vbaProject.bin", b"VBA DUMMY MACRO PROJECT")

    # 8. Synthetic standalone image
    img_path = output_dir / "sample_image.png"
    img = Image.new("RGB", (600, 200), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((30, 80), "LOCAL MARKDOWN OCR TEST SAMPLE 2026", fill=(0, 0, 0))
    img.save(img_path)

    # 9. Scanned PDF (single page image-only)
    scanned_pdf_path = output_dir / "scanned.pdf"
    scanned_canvas = canvas.Canvas(str(scanned_pdf_path), pagesize=letter)
    scanned_canvas.drawImage(str(img_path), 50, 500, width=500, height=160)
    scanned_canvas.showPage()
    scanned_canvas.save()

    # 10. Mixed PDF: Page 1 native, Page 2 scanned image, Page 3 blank, Page 4 native
    mixed_pdf_path = output_dir / "mixed.pdf"
    mc = canvas.Canvas(str(mixed_pdf_path), pagesize=letter)
    # Page 1: Native
    mc.setFont("Helvetica-Bold", 14)
    mc.drawString(72, 720, "Mixed Document - Page 1 Native")
    mc.setFont("Helvetica", 11)
    mc.drawString(72, 690, "Native digital text on page 1.")
    mc.showPage()
    # Page 2: Scanned image only (no native text)
    mc.drawImage(str(img_path), 50, 500, width=500, height=160)
    mc.showPage()
    # Page 3: Blank page
    mc.showPage()
    # Page 4: Native
    mc.setFont("Helvetica-Bold", 14)
    mc.drawString(72, 720, "Mixed Document - Page 4 Native Conclusion")
    mc.setFont("Helvetica", 11)
    mc.drawString(72, 690, "Native conclusion text on page 4.")
    mc.showPage()
    mc.save()

    print(f"Generated 10 fixtures successfully in: {output_dir}")


if __name__ == "__main__":
    fixtures_dir = Path(__file__).parent
    create_all_fixtures(fixtures_dir)
